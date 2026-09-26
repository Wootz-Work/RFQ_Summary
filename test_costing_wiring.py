"""
test_costing_wiring.py — the costing workbook in the product writeback.

The rules that matter: the extraction's rows reach the right tabs without any
bookkeeping columns, the uploaded file's id and link land in ttqlU / Vr8gz on
the RFQ row, and nothing in this path can cost an extraction that succeeded.

Run:
    python test_costing_wiring.py
"""
import io
import sys
import types

sys.path.insert(0, "src")

for name in ("fitz", "google", "google.oauth2", "google.oauth2.service_account",
             "google.api_core", "google.api_core.client_options",
             "google.cloud", "google.cloud.documentai_v1",
             "googleapiclient", "googleapiclient.discovery", "googleapiclient.http",
             "googleapiclient.errors"):
    sys.modules.setdefault(name, types.ModuleType(name))
sys.modules["google.oauth2.service_account"].Credentials = object
sys.modules["google.api_core.client_options"].ClientOptions = object
sys.modules["google.cloud.documentai_v1"].DocumentProcessorServiceClient = object
sys.modules["googleapiclient.discovery"].build = lambda *a, **k: None
sys.modules["googleapiclient.http"].MediaIoBaseDownload = object
sys.modules["googleapiclient.errors"].HttpError = Exception

import openpyxl

from rfq_summary import writer
from rfq_summary.config import Settings
from rfq_summary.costing_workbook import tabs_from_extraction
from rfq_summary.onedrive import UploadedFile
from rfq_summary.schema import (
    ExtractedProduct, ProductAnnexure, ProductExtractionHeader, ProductExtractionResult, TriageOutputPayload,
)

ok = True


def check(label, cond, detail=""):
    global ok
    ok &= bool(cond)
    print(("PASS  " if cond else "FAIL  ") + label + (f"  — {detail}" if detail and not cond else ""))


def _settings(**kw):
    base = dict(GLIDE_API_KEY="k", GLIDE_APP_ID="a",
                MS_GRAPH_TENANT_ID="t", MS_GRAPH_CLIENT_ID="c", MS_GRAPH_CLIENT_SECRET="s",
                EMAIL_FROM_ADDRESS="technology@wootz.work")
    base.update(kw)
    return Settings(**base)


COLS = ["variant_ref", "sr_no", "part_number", "description", "finish", "quantity", "target_price"]
ROWS = [[f"V{i}", str(i + 1), f"HB-{i:03d}", f"Hex bolt M{8 + 2 * i}", "HDG", f"{100 * (i + 1)} pcs", "0.40"]
        for i in range(5)]


def extraction():
    return ProductExtractionResult(
        header=ProductExtractionHeader(rfq_title="Project Falcon - Fasteners"),
        products=[
            ExtractedProduct(index=1, name="Hex Bolts", structure="family", variant_count=5,
                             annexure=ProductAnnexure(required=True, columns=COLS, rows=ROWS)),
            ExtractedProduct(index=2, name="Threaded Stud M56", structure="single", quantity="1,200 pcs",
                             details="- Stud M56 x 310\n- ASTM A193 B7", provenance={"quantity": "derived"}),
            ExtractedProduct(index=3, name="Their Washers", structure="family", quantity="As per annexure",
                             annexure=ProductAnnexure(required=True, by_reference=True)),
        ],
    )


# ---- extraction -> tabs ----------------------------------------------------
tabs = tabs_from_extraction(extraction())
check("one tab per family with rows, one for everything else", [t.name for t in tabs] == ["Hex Bolts", "Individual items"],
      str([t.name for t in tabs]))
fam = tabs[0]
headers = [h for h, _ in fam.columns]
check("no made-up reference or serial number columns", not any(h.lower().startswith(("variant", "sr")) for h in headers),
      str(headers))
check("part number and description shown in plain words", headers[:2] == ["Part number", "Description"], str(headers))
check("quantity is split into number and unit, not a column", "Quantity" not in headers
      and fam.lines[2].qty.value == 300 and fam.lines[2].qty_unit == "pcs")
check("each row is labelled for the Quotation by code and description", fam.lines[0].label == "HB-000 Hex bolt M8",
      fam.lines[0].label)
check("the family has one legend rate, named after it", fam.rate_groups == ["Hex Bolts"])
check("source lines taken from the variant count", fam.source_lines == 5)
ind = tabs[1]
check("single lines go to Individual items", [l.label for l in ind.lines] == ["Threaded Stud M56", "Their Washers"])
check("a derived quantity is red", ind.lines[0].qty.kind == "assume" and ind.lines[0].qty.value == 1200)
check("'As per annexure' leaves qty for the team", ind.lines[1].qty.kind == "input")
check("a by-reference family says why it is not expanded", "own workbook" in ind.lines[1].remarks)
check("weight is never set from the extraction", all(l.weight.kind == "input" for t in tabs for l in t.lines))


# ---- the writeback step ----------------------------------------------------
def run(settings=None, dest=("b!DRIVE", "01FOLDERIDAAAAAAAAAAAAAAAAAAAA"), upload=None, glide=None, template=None,
        ext=None):
    calls = {"upload": [], "glide": []}

    def fake_upload(s, drive_id, folder_id, filename, data):
        calls["upload"].append({"name": filename, "data": data, "folder": folder_id})
        if upload is not None:
            return upload(filename)
        return UploadedFile(id="01COSTINGFILEIDAAAAAAAAAAAAAAA", url="https://wootz-my.sharepoint.com/x/cost.xlsx",
                            name=filename)

    def fake_glide(s, row_id, values):
        calls["glide"].append((row_id, values))
        if glide is not None:
            return glide()
        return True

    saved = (writer.glide_fetch_annexure_destination, writer.upload_annexure, writer.glide_set_all_rfq_columns,
             writer._costing_template)
    writer.glide_fetch_annexure_destination = lambda s, r: dest
    writer.upload_annexure = fake_upload
    writer.glide_set_all_rfq_columns = fake_glide
    writer._costing_template = lambda s, run_id: template
    try:
        out = TriageOutputPayload(run_id="r1", row_id="RFQ1")
        calls["result"] = writer._attach_costing_workbook(settings or _settings(), out, "RFQ1", ext or extraction())
        return calls
    finally:
        (writer.glide_fetch_annexure_destination, writer.upload_annexure, writer.glide_set_all_rfq_columns,
         writer._costing_template) = saved


calls = run()
check("one workbook uploaded for the RFQ", len(calls["upload"]) == 1, str(calls["upload"]))
check("named after the RFQ", calls["upload"][0]["name"] == "Int Costing - Project Falcon - Fasteners.xlsx",
      calls["upload"][0]["name"])
check("file id to ttqlU and link to Vr8gz, on the RFQ row",
      calls["glide"] == [("RFQ1", {"ttqlU": "01COSTINGFILEIDAAAAAAAAAAAAAAA",
                                   "Vr8gz": "https://wootz-my.sharepoint.com/x/cost.xlsx"})], str(calls["glide"]))
check("reports success", calls["result"] is True)
book = openpyxl.load_workbook(io.BytesIO(calls["upload"][0]["data"]))
check("the uploaded file is the generated workbook",
      book.sheetnames == ["(Zai) Summary", "(Zai) Hex Bolts", "(Zai) Individual items"], str(book.sheetnames))

check("switched off means nothing happens", run(settings=_settings(ENABLE_COSTING_WORKBOOK="false"))["upload"] == [])
check("no graph credentials means nothing happens", run(settings=_settings(MS_GRAPH_CLIENT_SECRET=""))["upload"] == [])
check("no folder on the RFQ row means nothing is uploaded", run(dest=("b!DRIVE", ""))["upload"] == [])
c = run(upload=lambda n: None)
check("an upload that lands nowhere writes no link", c["glide"] == [] and c["result"] is False)


def _raise(*a, **k):
    raise RuntimeError("glide down")


c = run(glide=_raise)
check("a Glide failure after upload is contained", c["result"] is False and len(c["upload"]) == 1)
c = run(upload=_raise)
check("an upload that raises is contained", c["result"] is False and c["glide"] == [])
c = run(template=b"this is not a workbook")
check("a template we cannot open falls back to the generated tabs", len(c["upload"]) == 1 and c["result"] is True)
check("an extraction with no products uploads nothing",
      run(ext=ProductExtractionResult(header=ProductExtractionHeader(), products=[]))["upload"] == [])
check("the id and url columns can be renamed from config",
      run(settings=_settings(GLIDE_COL_ALL_RFQ_COSTING_FILE_ID="abc", GLIDE_COL_ALL_RFQ_COSTING_URL="xyz"))["glide"][0][1]
      .keys() == {"abc", "xyz"})

print("\nALL PASSED" if ok else "\nFAILURES ABOVE")
raise SystemExit(0 if ok else 1)
