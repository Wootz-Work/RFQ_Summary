"""
costing_workbook.py — the internal costing workbook, generated from what was
extracted from an RFQ.

One tab per product family, one tab for individual items, a summary tab so the
extraction can be checked against the source, and — when the team's template
is supplied — the template's own Quotation / Back-end / insight tabs, with the
Quotation's item table rewired to the generated tabs.

Rules the sheet follows, set by the costing team:
  * one row per product (per assembly for fabrication), nothing merged across rows
  * colour says where a value came from — orange fill: someone must fill it in;
    red text: our assumption; black: from the customer; pink: from the customer
    but technically doubtful; grey fill: formula
  * weight is never calculated — it is only ever what a drawing or the customer
    stated, otherwise an empty orange cell
  * no rate is ever filled in by us; rates live once in the legend under each
    table, so filling one cell prices every row that uses it
  * a number and its unit are separate columns
  * generated tabs are prefixed "(Zai)" so nobody mistakes them for typed work

Template round trip: openpyxl drops pictures and x14 data validations on load
(the template has a logo on three tabs and dropdowns pointing at Back-end), so
both are read straight from the template's zip and put back after the edit.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

ZAI_PREFIX = "(Zai) "
SUMMARY_TITLE = ZAI_PREFIX + "Summary"
PROCESS_SLOTS = 8
PROCESS_LIST_REACH = 200          # legend rows the formulas read; the list can grow this far
SPARE_PROCESS_ROWS = 12
FIRST_DATA_ROW = 3                # row 1 = column-group band, row 2 = headers

DEFAULT_PROCESSES = (
    "Cutting", "Rolling", "Bending", "Welding", "Machining", "Brushing", "Pickling & passivation",
    "Leak test", "Powder coating", "Painting", "Galvanising (HDG)", "Assembly",
)

# ----------------------------------------------------------------------------- values

KINDS = ("data", "assume", "doubt", "input", "calc")


@dataclass
class Val:
    """A cell value plus where it came from; the kind decides its colour."""
    value: Any = None
    kind: str = "data"

    def __post_init__(self):
        if self.kind not in KINDS:
            raise ValueError(f"unknown kind {self.kind!r}")
        if self.kind == "input":
            self.value = None        # an input cell is empty by definition


def data(v: Any) -> Val: return Val(v, "data")
def assumed(v: Any) -> Val: return Val(v, "assume")
def doubtful(v: Any) -> Val: return Val(v, "doubt")
def needed() -> Val: return Val(None, "input")


@dataclass
class Line:
    label: str                                   # item name carried onto the Quotation
    fields: Dict[str, Val]                       # customer-data columns, keyed by header
    qty: Val = field(default_factory=needed)
    qty_unit: str = "pcs"
    pack: Optional[Val] = None
    pack_unit: str = ""
    weight: Val = field(default_factory=needed)  # kg — stated, never computed
    rate_group: str = ""                         # weight_rate mode: which legend rate prices it
    material_rate: Val = field(default_factory=needed)   # process mode
    processes: List[Val] = field(default_factory=list)   # process mode: names from the legend
    remarks: str = ""


@dataclass
class Tab:
    name: str                                    # "Washers" -> sheet "(Zai) Washers"
    kind: str                                    # "family" | "individual"
    mode: str                                    # "weight_rate" | "process"
    columns: List[Tuple[str, int]]               # customer-data columns: (header, width)
    lines: List[Line]
    source: str = ""                             # where it came from, e.g. "Washers.xlsx › Washers"
    source_lines: Optional[int] = None           # data lines counted in the source
    skipped: List[Tuple[str, str]] = field(default_factory=list)   # (source row, reason)
    rate_groups: List[str] = field(default_factory=list)            # weight_rate legend rates
    processes: Tuple[str, ...] = DEFAULT_PROCESSES                  # process legend

    @property
    def title(self) -> str:
        return (ZAI_PREFIX + self.name)[:31]


@dataclass
class Commons:
    """Assumptions shared by every tab. Set once, on the first generated tab."""
    currency: str = "GBP"
    fx: float = 120.0                     # INR per 1 unit of currency
    packaging: float = 0.02
    margin: float = 0.20
    pallet_capacity: float = 950.0        # kg / pallet
    price_per_pallet: float = 50000.0     # INR / pallet
    basis: str = "Template default"


# ----------------------------------------------------------------------------- styling

_BLACK, _RED, _PINK = "000000", "C00000", "E0218A"
_ORANGE = PatternFill("solid", fgColor="FCD5B4")
_GREY = PatternFill("solid", fgColor="F2F2F2")
_NAVY = PatternFill("solid", fgColor="1F3864")
_BAND = {"cust": "D9E1F2", "wt": "E2EFDA", "proc": "E2EFDA", "cost": "FFF2CC", "rem": "EDEDED"}
_BAND_NAME = {"cust": "CUSTOMER DATA", "wt": "WEIGHT & RATE",
              "proc": "PROCESSES — pick from the list; clear a cell to remove",
              "cost": "COST BUILD-UP  (formulas)", "rem": ""}
_thin = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _font(kind: str, bold: bool = False, size: int = 10) -> Font:
    colour = {"assume": _RED, "doubt": _PINK}.get(kind, _BLACK)
    return Font(name="Calibri", size=size, bold=bold, color=colour)


def _paint(cell, val: Val, number_format: Optional[str] = None) -> None:
    if val.kind == "input":
        cell.value = None
        cell.fill = _ORANGE
    else:
        cell.value = val.value
        cell.font = _font(val.kind)
        if val.kind == "calc":
            cell.fill = _GREY
    if number_format:
        cell.number_format = number_format


def _formula(cell, text: str, number_format: Optional[str] = None, bold: bool = False) -> None:
    cell.value = text
    cell.font = _font("calc", bold)
    cell.fill = _GREY
    if number_format:
        cell.number_format = number_format


# ----------------------------------------------------------------------------- one tab

@dataclass
class _Built:
    """What the Quotation and the summary need to know about a built tab."""
    title: str
    rows: List[int]
    col: Dict[str, str]
    counts: Dict[str, int]


def _cost_columns(currency: str) -> List[Tuple[str, int, str]]:
    return [("Mfg cost / pc (INR)", 10, "0.000"), ("Ex-works / pc (INR)", 10, "0.000"),
            ("Freight / pc (INR)", 9, "0.000"), ("DAP / pc (INR)", 9, "0.000"),
            (f"Ex-works / pc ({currency})", 10, "0.0000"), (f"DAP / pc ({currency})", 10, "0.0000"),
            (f"Annual DAP value ({currency})", 12, "#,##0")]


def _layout(tab: Tab, currency: str) -> List[Tuple[str, str, int]]:
    has_pack = any(l.pack is not None for l in tab.lines)
    cols = [(h, "cust", w) for h, w in tab.columns]
    cols += [("Annual qty", "cust", 10), ("Qty unit", "cust", 7)]
    if has_pack:
        cols += [("Pack size", "cust", 8), ("Pack unit", "cust", 9)]
    cols += [("Weight", "wt", 9), ("Weight unit", "wt", 7)]
    if tab.mode == "weight_rate":
        cols += [("Rate (INR / kg)", "wt", 9)]
    else:
        cols += [("Material rate (INR / kg)", "wt", 10)]
        cols += [(f"Process {i}", "proc", 13) for i in range(1, PROCESS_SLOTS + 1)]
        cols += [("Bought-out items / pc (INR)", "proc", 11)]
    cols += [("Missing input", "wt" if tab.mode == "weight_rate" else "proc", 18)]
    cols += [(h, "cost", w) for h, w, _ in _cost_columns(currency)]
    cols += [("Remarks", "rem", 46)]
    return cols


def _build_tab(wb, tab: Tab, commons: Commons, shared: Dict[str, str], first_tab: bool) -> _Built:
    if tab.mode == "process":
        used = [p.value for l in tab.lines for p in l.processes if p.value]
        extra = [p for p in dict.fromkeys(used) if p not in tab.processes]
        tab = Tab(**{**tab.__dict__, "processes": tuple(tab.processes) + tuple(extra)})
    ws = wb.create_sheet(tab.title)
    cur = commons.currency
    cols = _layout(tab, cur)
    col = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(cols)}
    n = len(tab.lines)
    first, last = FIRST_DATA_ROW, FIRST_DATA_ROW + max(n, 1) - 1

    # ---- the legend's position is fixed before any formula is written
    key_top = last + 3
    a_head = key_top + 7
    common_rows = [
        (f"INR per {cur}", commons.fx, f"INR / {cur}"),
        ("Packaging", commons.packaging, "%"),
        ("Margin", commons.margin, "%"),
        ("Pallet capacity", commons.pallet_capacity, "kg / pallet"),
        ("Price per pallet", commons.price_per_pallet, "INR / pallet"),
    ]
    ref: Dict[str, str] = {}
    r = a_head + 2
    for name, *_ in common_rows:
        ref[name] = f"$C${r}"
        r += 1
    rate_rows: Dict[str, int] = {}
    if tab.mode == "weight_rate":
        for g in tab.rate_groups:
            rate_rows[g] = r
            ref[g] = f"$C${r}"
            r += 1
        proc_first = None
    else:
        proc_first = r
        names = f"$B${proc_first}:$B${proc_first + PROCESS_LIST_REACH - 1}"
        rates = f"$C${proc_first}:$C${proc_first + PROCESS_LIST_REACH - 1}"

    # ---- band + headers
    spans: Dict[str, List[int]] = {}
    for i, (_, band, _) in enumerate(cols, 1):
        spans.setdefault(band, [i, i])[1] = i
    for band, (a, b) in spans.items():
        c = ws.cell(1, a, _BAND_NAME[band])
        c.font = Font(size=9, bold=True, color="1F3864")
        c.alignment = Alignment(horizontal="center")
        for j in range(a, b + 1):
            ws.cell(1, j).fill = PatternFill("solid", fgColor=_BAND[band])
        if b > a:
            ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
    for i, (h, _, w) in enumerate(cols, 1):
        c = ws.cell(2, i, h)
        c.font = Font(size=9, bold=True, color="FFFFFF")
        c.fill = _NAVY
        c.border = _BOX
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 40
    cost_fmt = {h: f for h, _, f in _cost_columns(cur)}
    widths = {h: w for h, _, w in cols}

    counts = {"doubt": 0, "assume": 0, "input": 0}
    rows_out: List[int] = []
    for k, line in enumerate(tab.lines):
        r = first + k
        rows_out.append(r)
        at = lambda h: ws[f"{col[h]}{r}"]
        for h, *_ in cols:
            c = at(h)
            c.border = _BOX
            c.alignment = Alignment(vertical="center", wrap_text=(h == "Remarks" or widths.get(h, 0) >= 40))

        def put(h: str, v: Val, fmt: Optional[str] = None):
            _paint(at(h), v, fmt)
            if v.kind in counts:
                counts[v.kind] += 1

        for h, _w in tab.columns:
            put(h, line.fields.get(h, data(None)))
        put("Annual qty", line.qty, "#,##0")
        at("Qty unit").value = line.qty_unit
        at("Qty unit").font = _font("data")
        if "Pack size" in col and line.pack is not None:
            put("Pack size", line.pack, "#,##0")
            at("Pack unit").value = line.pack_unit
            at("Pack unit").font = _font("data")
        put("Weight", line.weight, "0.000##")
        at("Weight unit").value = "kg"
        at("Weight unit").font = _font("data")
        w = f"{col['Weight']}{r}"

        if tab.mode == "weight_rate":
            g = ref.get(line.rate_group)
            rate = f"{col['Rate (INR / kg)']}{r}"
            if g:
                _formula(at("Rate (INR / kg)"), f'=IF({g}="","",{g})', "0")
                group = line.rate_group.replace('"', "'")
                missing = f'=IF({w}="","Add weight",IF({rate}="","Add rate: {group}",""))'
            else:
                put("Rate (INR / kg)", needed())
                missing = f'=IF({w}="","Add weight",IF({rate}="","Add rate",""))'
            mfg = f'=IF(OR({w}="",{rate}=""),"",{w}*{rate})'
        else:
            mat = f"{col['Material rate (INR / kg)']}{r}"
            put("Material rate (INR / kg)", line.material_rate, "0")
            for i in range(PROCESS_SLOTS):
                slot = at(f"Process {i + 1}")
                if i < len(line.processes):
                    put(f"Process {i + 1}", line.processes[i])
                slot.alignment = Alignment(horizontal="center", vertical="center")
            put("Bought-out items / pc (INR)", needed(), "0.00")
            p = f"{col['Process 1']}{r}:{col[f'Process {PROCESS_SLOTS}']}{r}"
            unpriced = f'SUMPRODUCT(({p}<>"")*(COUNTIFS({names},{p},{rates},"<>")=0))'
            # A process without a rate is left out of the total, not allowed to blank it —
            # the Missing-input column says how many are still unpriced.
            mfg = (f'=IF(OR({w}="",{mat}=""),"",{w}*({mat}+SUMPRODUCT(SUMIF({names},{p},{rates})))'
                   f'+N({col["Bought-out items / pc (INR)"]}{r}))')
            missing = (f'=IF({w}="","Add weight",IF({mat}="","Add material rate",'
                       f'IF({unpriced}>0,{unpriced}&" process rate(s) missing — not in total","")))')
        _formula(at("Missing input"), missing)

        m, e, fr, d = (f"{col[h]}{r}" for h in ("Mfg cost / pc (INR)", "Ex-works / pc (INR)",
                                                 "Freight / pc (INR)", "DAP / pc (INR)"))
        qty = f"{col['Annual qty']}{r}"
        dap_cur = f"{col[f'DAP / pc ({cur})']}{r}"
        chain = {
            "Mfg cost / pc (INR)": mfg,
            "Ex-works / pc (INR)": f'=IF({m}="","",{m}*(1+{ref["Packaging"]})/(1-{ref["Margin"]}))',
            "Freight / pc (INR)": f'=IF({w}="","",{w}*{ref["Price per pallet"]}/{ref["Pallet capacity"]})',
            "DAP / pc (INR)": f'=IF(OR({e}="",{fr}=""),"",{e}+{fr})',
            f"Ex-works / pc ({cur})": f'=IF({e}="","",{e}/{ref[f"INR per {cur}"]})',
            f"DAP / pc ({cur})": f'=IF({d}="","",{d}/{ref[f"INR per {cur}"]})',
            f"Annual DAP value ({cur})": f'=IF(OR({dap_cur}="",{qty}=""),"",{dap_cur}*{qty})',
        }
        for h, f in chain.items():
            _formula(at(h), f, cost_fmt[h])
        at("Remarks").value = line.remarks or None
        at("Remarks").font = _font("data")

    ws.freeze_panes = ws.cell(first, 3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:2"

    # ---- legend
    ws.cell(key_top, 2, "LEGEND").font = Font(size=12, bold=True, color="1F3864")
    key = [("", "input", "Orange — to fill in (from the drawing, the customer, or your rate). Left empty on purpose"),
           ("Red text", "assume", "Our assumption — fair, kept on the higher side. Check before quoting"),
           ("Black text", "data", "Straight from the customer's data or drawing"),
           ("Pink text", "doubt", "In the customer's data, but technically doubtful — confirm"),
           ("Grey cell", "calc", "Formula — fills in once its inputs are there; don't type over it")]
    for i, (txt, kind, expl) in enumerate(key):
        c = ws.cell(key_top + 1 + i, 2, txt or None)
        c.border = _BOX
        if kind == "input":
            c.fill = _ORANGE
        else:
            c.font = _font(kind, True)
            if kind == "calc":
                c.fill = _GREY
        ws.cell(key_top + 1 + i, 3, expl).font = Font(size=10)
    ws.cell(a_head, 2, "ASSUMPTIONS & RATES — change once here; every row updates").font = \
        Font(size=11, bold=True, color="1F3864")
    for c_, t in ((2, "Item"), (3, "Value"), (4, "Unit"), (5, "Basis")):
        ws.cell(a_head + 1, c_, t).font = Font(size=9, bold=True, color="595959")

    def legend_row(r: int, name: str, value: Any, unit: str, basis: str, kind: str, bold_name=False):
        nc = ws.cell(r, 2, name or None)
        nc.font = Font(size=10, bold=bold_name)
        if bold_name:
            nc.border = _BOX
            if not name:
                nc.fill = _ORANGE
        c = ws.cell(r, 3)
        c.border = _BOX
        _paint(c, Val(value, kind), "0%" if unit == "%" else None)
        ws.cell(r, 4, unit).font = Font(size=10, color="595959")
        ws.cell(r, 5, basis).font = Font(size=9, italic=True, color="595959")

    r = a_head + 2
    for name, value, unit in common_rows:
        if first_tab:
            legend_row(r, name, value, unit, commons.basis, "assume")
            shared[name] = f"'{tab.title}'!$C${r}"
        else:
            legend_row(r, name, f"={shared[name]}", unit, "Linked to the first tab — change it there", "calc")
        r += 1
    if tab.mode == "weight_rate":
        for g in tab.rate_groups:
            legend_row(rate_rows[g], g, None, "INR / kg", "Your rate — prices every row of this group", "input")
    else:
        for i, p in enumerate(tab.processes):
            legend_row(proc_first + i, p, None, "INR / kg", "Your rate — added when a row lists this process",
                       "input", bold_name=True)
        for j in range(SPARE_PROCESS_ROWS):
            legend_row(proc_first + len(tab.processes) + j, "", None, "INR / kg",
                       "Spare — type a new process here. Need more? keep going on the rows below",
                       "input", bold_name=True)
        # the dropdown grows with the list: it counts the names typed so far
        dn = f"ZaiProcesses{len(wb.defined_names) + 1}"
        sheet = f"'{tab.title}'"
        wb.defined_names[dn] = DefinedName(
            dn, attr_text=f"OFFSET({sheet}!$B${proc_first},0,0,"
                          f"MAX(1,COUNTA({sheet}!$B${proc_first}:$B${proc_first + PROCESS_LIST_REACH - 1})),1)")
        dv = DataValidation(type="list", formula1=dn, allow_blank=True, showErrorMessage=True,
                            error="Pick a process from the list, or add it to the legend first")
        ws.add_data_validation(dv)
        dv.add(f"{col['Process 1']}{first}:{col[f'Process {PROCESS_SLOTS}']}{last + 50}")

    return _Built(tab.title, rows_out, col, counts)


# ----------------------------------------------------------------------------- summary

def _build_summary(wb, tabs: List[Tab], built: List[_Built], currency: str) -> str:
    """Tab-by-tab reconciliation against the source, then an index of every line."""
    ws = wb.create_sheet(SUMMARY_TITLE)
    head = ["Tab", "Type", "Source", "Lines in source", "Lines extracted", "Not extracted",
            "Doubtful (pink)", "Assumed (red)", "To fill in (orange)", "Check"]
    widths = [24, 12, 36, 10, 10, 10, 10, 10, 10, 34]

    def header_row(r: int, labels: List[str]):
        for i, h in enumerate(labels, 1):
            c = ws.cell(r, i, h)
            c.font = Font(size=9, bold=True, color="FFFFFF")
            c.fill = _NAVY
            c.border = _BOX
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 32

    header_row(1, head)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for t, b in zip(tabs, built):
        src = t.source_lines
        ext = len(t.lines)
        skipped = (src - ext) if src is not None else None
        check = ("All source lines extracted" if skipped == 0 else
                 f"{skipped} line(s) not extracted — see below" if skipped else "Count the source to confirm")
        vals = [t.title, "Family" if t.kind == "family" else "Individual items", t.source, src, ext, skipped,
                b.counts["doubt"], b.counts["assume"], b.counts["input"], check]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = _BOX
            c.font = _font("doubt" if (i == 10 and skipped) else "data")
        ws.cell(r, 1).hyperlink = f"#'{t.title}'!A1"
        ws.cell(r, 1).font = Font(size=10, color="1F3864", underline="single")
        r += 1
    total = ["Total", f"{sum(t.kind == 'family' for t in tabs)} families, "
                      f"{sum(len(t.lines) for t in tabs if t.kind == 'individual')} individual items", "",
             sum(t.source_lines or 0 for t in tabs), sum(len(t.lines) for t in tabs),
             sum(((t.source_lines or len(t.lines)) - len(t.lines)) for t in tabs),
             sum(b.counts["doubt"] for b in built), sum(b.counts["assume"] for b in built),
             sum(b.counts["input"] for b in built), ""]
    for i, v in enumerate(total, 1):
        c = ws.cell(r, i, v)
        c.font = Font(size=10, bold=True)
        c.border = _BOX
        c.fill = _GREY

    skipped_rows = [(t.title, s, why) for t in tabs for s, why in t.skipped]
    r += 3
    ws.cell(r - 1, 1, "Not extracted from the source").font = Font(size=11, bold=True, color="1F3864")
    header_row(r, ["Tab", "Source row", "Reason"])
    r += 1
    if not skipped_rows:
        ws.cell(r, 1, "Nothing — every source line is in a tab").font = Font(size=10, italic=True)
        r += 1
    for title, s, why in skipped_rows:
        for i, v in enumerate((title, s, why), 1):
            ws.cell(r, i, v).border = _BOX
        r += 1

    r += 2
    ws.cell(r - 1, 1, "Every extracted line").font = Font(size=11, bold=True, color="1F3864")
    header_row(r, ["Tab", "#", "Item", "Annual qty", "Qty unit", f"DAP / pc ({currency})", "Go to row"])
    r += 1
    item_first = r
    for t, b in zip(tabs, built):
        dap = b.col[f"DAP / pc ({currency})"]
        for k, (line, row) in enumerate(zip(t.lines, b.rows), 1):
            sheet = f"'{t.title}'"
            vals = [t.title, k, line.label, f"={sheet}!{b.col['Annual qty']}{row}",
                    f"={sheet}!{b.col['Qty unit']}{row}", f'=IF({sheet}!{dap}{row}="","",{sheet}!{dap}{row})',
                    f"Row {row}"]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.border = _BOX
                c.font = _font("calc" if isinstance(v, str) and v.startswith("=") else "data")
            ws.cell(r, 4).number_format = "#,##0"
            ws.cell(r, 6).number_format = "0.0000"
            ws.cell(r, 7).hyperlink = f"#{sheet}!A{row}"
            ws.cell(r, 7).font = Font(size=10, color="1F3864", underline="single")
            r += 1
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    items =f"'{SUMMARY_TITLE}'!$C${item_first}:$C${max(item_first, r - 1)}"
    wb.defined_names["ZaiItems"] = DefinedName("ZaiItems", attr_text=items)
    return "ZaiItems"


# ----------------------------------------------------------------------------- template surgery

def shift_formula(formula: str, start: int, k: int) -> str:
    """Move every same-sheet reference at or below `start` down by `k` rows.

    Rows are inserted into the Quotation above its freight/terms block; every
    formula that points into that block must follow it, and nothing else may
    move — the template's lookups into 'Back-end'!F2:K6 and its relative
    reference to A4 must stay exactly where they are.
    """
    if not isinstance(formula, str) or not formula.startswith("=") or k == 0:
        return formula

    def bump(m: re.Match) -> str:
        row = int(m.group(3))
        return f"{m.group(1)}{m.group(2)}{row + k if row >= start else row}"

    tok = Tokenizer(formula)
    out = []
    for t in tok.items:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" not in v:
            v = re.sub(r"(\$?[A-Z]{1,3})(\$?)(\d+)", bump, v)
        out.append(v)
    return "=" + "".join(out)


def _shift_sqref(sqref: str, start: int, k: int) -> str:
    parts = []
    for rng in str(sqref).split():
        parts.append(re.sub(r"([A-Z]{1,3})(\d+)",
                            lambda m: f"{m.group(1)}{int(m.group(2)) + k if int(m.group(2)) >= start else m.group(2)}",
                            rng))
    return " ".join(parts)


@dataclass
class _Picture:
    sheet: str
    anchor_xml: str
    data: bytes
    ext: str


def _relationships(xml: str) -> Dict[str, str]:
    """Id -> Target for every relationship, whatever order the attributes come in."""
    out = {}
    for rel in re.findall(r"<Relationship\b[^>]*>", xml):
        rid, target = re.search(r'\bId="([^"]+)"', rel), re.search(r'\bTarget="([^"]+)"', rel)
        if rid and target:
            out[rid.group(1)] = target.group(1)
    return out


def _part(target: str, base: str) -> str:
    """Resolve a relationship target to a zip path ("worksheets/sheet1.xml" -> "xl/worksheets/sheet1.xml")."""
    t = target.lstrip("/")
    return t if t.startswith("xl/") else base + t


def _template_parts(template: bytes) -> Tuple[List[_Picture], Dict[str, List[Tuple[str, str]]]]:
    """Pictures and x14 dropdowns, per sheet name — both are lost by openpyxl."""
    z = zipfile.ZipFile(io.BytesIO(template))
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    rid_target = _relationships(z.read("xl/_rels/workbook.xml.rels").decode("utf-8"))
    pictures: List[_Picture] = []
    dvs: Dict[str, List[Tuple[str, str]]] = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        name = name.replace("&amp;", "&")
        path = _part(rid_target[rid], "xl/")
        sheet_xml = z.read(path).decode("utf-8")
        for f, sq in re.findall(r"<x14:dataValidation[^>]*>.*?<xm:f>(.*?)</xm:f>.*?<xm:sqref>(.*?)</xm:sqref>",
                                sheet_xml, re.S):
            dvs.setdefault(name, []).append((f.replace("&amp;", "&").replace("&apos;", "'"), sq))
        rel_path = path.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rel_path not in z.namelist():
            continue
        for dtarget in re.findall(r'Target="([^"]*drawings/[^"]+)"', z.read(rel_path).decode("utf-8")):
            dpath = "xl/drawings/" + dtarget.split("drawings/")[-1]
            # Excel writes <xdr:from>, other tools write <from>; strip tag prefixes so both read the same
            dxml = re.sub(r"<(/?)[A-Za-z]\w*:", r"<\1", z.read(dpath).decode("utf-8"))
            drels_path = dpath.replace("drawings/", "drawings/_rels/") + ".rels"
            drels = z.read(drels_path).decode("utf-8") if drels_path in z.namelist() else ""
            media = _relationships(drels)
            for anchor in re.findall(r"<(?:twoCellAnchor|oneCellAnchor)\b.*?</(?:twoCellAnchor|oneCellAnchor)>",
                                     dxml, re.S):
                emb = re.search(r'r:embed="([^"]+)"', anchor)
                if not emb or emb.group(1) not in media:
                    continue
                mpath = "xl/media/" + media[emb.group(1)].split("media/")[-1]
                pictures.append(_Picture(name, anchor, z.read(mpath), mpath.rsplit(".", 1)[-1]))
    return pictures, dvs


def _add_picture(ws, pic: _Picture, row_start: int, k: int) -> None:
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    def marker(tag: str) -> Optional[AnchorMarker]:
        m = re.search(rf"<{tag}>\s*<col>(\d+)</col>\s*<colOff>(-?\d+)</colOff>\s*"
                      rf"<row>(\d+)</row>\s*<rowOff>(-?\d+)</rowOff>\s*</{tag}>", pic.anchor_xml)
        if not m:
            return None
        col, col_off, row, row_off = map(int, m.groups())
        if row + 1 >= row_start:
            row += k
        return AnchorMarker(col=col, colOff=col_off, row=row, rowOff=row_off)

    img = XLImage(io.BytesIO(pic.data))
    frm, to = marker("from"), marker("to")
    ext = re.search(r'<ext cx="(\d+)" cy="(\d+)"', pic.anchor_xml)
    if pic.anchor_xml.startswith("<twoCellAnchor") and frm and to:
        edit = re.search(r'editAs="(\w+)"', pic.anchor_xml)
        img.anchor = TwoCellAnchor(editAs=edit.group(1) if edit else None, _from=frm, to=to)
    elif frm and ext:
        img.anchor = OneCellAnchor(_from=frm, ext=XDRPositiveSize2D(int(ext.group(1)), int(ext.group(2))))
    else:
        return
    ws.add_image(img)


def _find_row(ws, text: str, after: int = 1, column: int = 1) -> Optional[int]:
    for r in range(after, ws.max_row + 1):
        v = ws.cell(r, column).value
        if isinstance(v, str) and v.strip().lower().startswith(text.lower()):
            return r
    return None


def _shift_down(ws, start: int, k: int) -> None:
    """Insert k rows above `start`, carrying formulas, merges, heights and dropdowns."""
    max_col = get_column_letter(ws.max_column)
    merged = [str(m) for m in ws.merged_cells.ranges]
    below = [m for m in merged if range_boundaries(m)[1] >= start]
    for m in below:
        ws.unmerge_cells(m)
    heights = {r: ws.row_dimensions[r].height for r in range(start, ws.max_row + 1)
               if ws.row_dimensions[r].height is not None}
    ws.move_range(f"A{start}:{max_col}{ws.max_row}", rows=k, translate=False)
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = shift_formula(c.value, start, k)
    for m in below:
        c1, r1, c2, r2 = range_boundaries(m)
        ws.merge_cells(start_row=r1 + k, start_column=c1, end_row=r2 + k, end_column=c2)
    for r in range(start, start + k):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[r + k].height = h
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_shift_sqref(str(dv.sqref), start, k))
    # Conditional formats are left alone: the template's only rules (B7, B10) sit above the table.


def _wire_quotation(ws, entries: List[Tuple[str, str, int, Dict[str, str]]], currency: str,
                    template_dvs: List[Tuple[str, str]], items_name: str) -> Tuple[int, int]:
    """Point the Quotation's item table at the generated tabs. Returns (start, rows inserted)."""
    head = _find_row(ws, "Item Name")
    if head is None:
        return 0, 0
    first = head + 1
    freight = _find_row(ws, "Freight", first) or (first + 12)
    capacity = freight - first
    n = len(entries)
    k = max(0, n - capacity)
    if k:
        _shift_down(ws, freight, k)
    template_row = first + 1 if capacity > 1 else first
    styles = {c: ws.cell(template_row, c)._style for c in range(1, 6)}
    for i in range(max(n, capacity)):
        r = first + i
        for c in range(1, 6):
            cell = ws.cell(r, c)
            cell.value = None
            if i >= capacity:
                cell._style = styles[c]
        if i >= n:
            continue
        label, sheet, row, col = entries[i]
        s = f"'{sheet}'"
        exw = f"{s}!{col[f'Ex-works / pc ({currency})']}{row}"
        ws.cell(r, 1).value = label
        ws.cell(r, 2).value = f'=IF({exw}="","",{exw}*100)'
        qty = f"{s}!{col['Annual qty']}{row}"
        ws.cell(r, 3).value = f'=IF({qty}="","",{qty})'
        ws.cell(r, 4).value = f'=IF(B{r}="","",C{r}*B{r}/100)'
    last = first + max(n, 1) - 1
    total = _find_row(ws, "Total", freight + k)
    if total:
        # Item prices are ex-works; the template's Freight line is what turns the sum into DAP.
        ws.cell(total, 4).value = f"=SUM(D{first}:D{last})+N(D{freight + k})"

    for f, sq in template_dvs:
        if "Costing!" in f:
            f, sq = items_name, f"A{first}:A{last}"
        else:
            sq = _shift_sqref(sq, freight, k)
        dv = DataValidation(type="list", formula1=f, allow_blank=True)
        ws.add_data_validation(dv)
        dv.sqref = MultiCellRange(sq)
    return freight, k


# ----------------------------------------------------------------------------- entry point

def build_costing_workbook(tabs: List[Tab], *, template: Optional[bytes] = None,
                           commons: Optional[Commons] = None) -> bytes:
    """Build the workbook. With a template, its tabs are kept and the Quotation rewired."""
    commons = commons or Commons()
    if not tabs:
        raise ValueError("nothing to write — no tabs")
    pictures: List[_Picture] = []
    x14: Dict[str, List[Tuple[str, str]]] = {}
    if template:
        pictures, x14 = _template_parts(template)
        wb = openpyxl.load_workbook(io.BytesIO(template))
        if "Costing" in wb.sheetnames:
            del wb["Costing"]
        insert_at = wb.sheetnames.index("Back-end") + 1 if "Back-end" in wb.sheetnames else len(wb.sheetnames)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        insert_at = 0

    shared: Dict[str, str] = {}
    built = [_build_tab(wb, t, commons, shared, first_tab=(i == 0)) for i, t in enumerate(tabs)]
    items_name = _build_summary(wb, tabs, built, commons.currency)

    order = [SUMMARY_TITLE] + [b.title for b in built]
    for pos, title in enumerate(order):
        sheet = wb[title]
        wb._sheets.remove(sheet)
        wb._sheets.insert(insert_at + pos, sheet)

    if template:
        shifts: Dict[str, Tuple[int, int]] = {}
        if "Quotation" in wb.sheetnames:
            entries = [(line.label, b.title, row, b.col)
                       for t, b in zip(tabs, built) for line, row in zip(t.lines, b.rows)]
            shifts["Quotation"] = _wire_quotation(wb["Quotation"], entries, commons.currency,
                                                  x14.get("Quotation", []), items_name)
        for sheet, rules in x14.items():
            if sheet in ("Quotation", "Costing") or sheet not in wb.sheetnames:
                continue
            for f, sq in rules:
                dv = DataValidation(type="list", formula1=items_name if "Costing!" in f else f, allow_blank=True)
                wb[sheet].add_data_validation(dv)
                dv.sqref = MultiCellRange(sq)
        # openpyxl keeps pictures only when Pillow is installed, and never x14 metadata; drop
        # whatever it kept and restore every picture from the template, so both cases end the same.
        for sheet in {p.sheet for p in pictures}:
            if sheet in wb.sheetnames:
                wb[sheet]._images = []
        for pic in pictures:
            if pic.sheet in wb.sheetnames:
                start, k = shifts.get(pic.sheet, (0, 0))
                _add_picture(wb[pic.sheet], pic, start or 10 ** 9, k)
        wb.active = wb.sheetnames.index("Quotation") if "Quotation" in wb.sheetnames else 0
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == wb.active.title

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------- from an extraction

_QTY_HEADER = re.compile(r"^\s*(qty|quantity)\b|annual\s*(qty|quantity)|\busage\b", re.IGNORECASE)
_SHEET_ILLEGAL = re.compile(r"[\[\]:*?/\\]")


def split_quantity(text: Any) -> Tuple[Val, str]:
    """'1,200 pcs (annual)' -> (1200, 'pcs'). No number means someone has to fill it in."""
    if isinstance(text, (int, float)) and not isinstance(text, bool):
        return data(text), "pcs"
    s = str(text or "").strip()
    m = re.search(r"\d[\d,]*(?:\.\d+)?", s)
    if not m:
        return needed(), ""
    n = float(m.group().replace(",", ""))
    unit = re.match(r"\s*([A-Za-z][A-Za-z/.]*)", s[m.end():])
    return data(int(n) if n.is_integer() else n), (unit.group(1).lower().rstrip(".") if unit else "pcs")


def _tab_name(name: str, taken: set) -> str:
    base = _SHEET_ILLEGAL.sub(" ", str(name or "").strip()) or "Family"
    base = re.sub(r"\s+", " ", base)[:25].strip()
    out, i = base, 2
    while out.lower() in taken:
        out = f"{base[:22]} {i}"
        i += 1
    taken.add(out.lower())
    return out


def _width(header: str, values: List[Any], cap: int = 40) -> int:
    return max(8, min(cap, max([len(str(header))] + [len(str(v or "")) for v in values]) + 2))


def tabs_from_extraction(extraction: Any) -> List[Tab]:
    """
    Turn a product extraction into workbook tabs: one per family, one for the rest.

    Only what the extraction actually carries is used. Family rows come across
    column for column, minus bookkeeping (serial numbers, made-up references);
    the product-level quantity turns red when the extraction marked it derived.
    Weight and rates are never set here — they are the team's to fill.
    """
    from .quote_sheet import _annexure_rows_as_dicts
    from .sheet_columns import display_header, visible_columns

    tabs: List[Tab] = []
    taken: set = set()
    individual: List[Line] = []
    has_drawing = False
    for p in getattr(extraction, "products", None) or []:
        name = str(getattr(p, "name", "") or "").strip()
        structure = str(getattr(p, "structure", "") or "").strip().lower()
        annexure = getattr(p, "annexure", None)
        rows = _annexure_rows_as_dicts(annexure) if annexure else []
        if structure == "family" and rows:
            cols = visible_columns(getattr(annexure, "columns", None) or [], rows)
            qty_col = next((c for c in cols if _QTY_HEADER.search(display_header(c))), None)
            shown = [c for c in cols if c != qty_col]
            label_col = next((c for c in shown if display_header(c).lower() in ("description", "part name")),
                             shown[0] if shown else None)
            code_col = next((c for c in shown if display_header(c).lower() in ("part number", "item code", "stock code")), None)
            lines = []
            for row in rows:
                qty, unit = split_quantity(row.get(qty_col)) if qty_col else (needed(), "pcs")
                label = str(row.get(label_col) or "").strip() if label_col else ""
                if code_col and code_col != label_col and row.get(code_col):
                    label = f"{row.get(code_col)} {label}".strip()
                lines.append(Line(label=label or name, fields={display_header(c): data(row.get(c) or None) for c in shown},
                                  qty=qty, qty_unit=unit or "pcs", rate_group=name or "This family"))
            tabs.append(Tab(
                name=_tab_name(name, taken), kind="family", mode="weight_rate",
                columns=[(display_header(c), _width(display_header(c), [r.get(c) for r in rows])) for c in shown],
                lines=lines,
                source=" · ".join(x for x in (f"Line {p.index}" if getattr(p, "index", None) is not None else "",
                                              str(getattr(p, "source_ref", "") or "").strip()) if x),
                source_lines=getattr(p, "variant_count", None) or len(rows),
                rate_groups=[name or "This family"],
            ))
            continue

        qty, unit = split_quantity(getattr(p, "quantity", ""))
        if qty.kind == "data" and str((getattr(p, "provenance", None) or {}).get("quantity", "")).lower() == "derived":
            qty = assumed(qty.value)
        dwg = str(getattr(p, "dwg_link", "") or "").strip()
        has_drawing = has_drawing or bool(dwg)
        remarks = ("Variants are in the customer's own workbook — not expanded here"
                   if structure == "family" else "")
        individual.append(Line(label=name, fields={"Part name": data(name or None),
                                                   "Details": data(str(getattr(p, "details", "") or "").strip() or None),
                                                   "Drawing": data(dwg or None)},
                               qty=qty, qty_unit=unit or "pcs", remarks=remarks))
    if individual:
        cols = [("Part name", 30), ("Details", 60)] + ([("Drawing", 24)] if has_drawing else [])
        tabs.append(Tab(name="Individual items", kind="individual", mode="process", columns=cols,
                        lines=individual, source="Line items in the RFQ", source_lines=len(individual)))
    return tabs
