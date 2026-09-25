"""
test_quote_sheet.py — the supplier quotation sheet for a product family.

What matters here: an attribute identical on every row never becomes a
column, the customer's own workbook is read in full rather than retyped by a
model, and nothing that reaches the supplier identifies the customer or
names a target price.

Run:
    python test_quote_sheet.py
"""
import io
import sys
import time

sys.path.insert(0, "src")

from openpyxl import Workbook, load_workbook

from rfq_summary.quote_sheet import (
    DEFAULT_QUOTE_COLUMNS,
    QuoteColumn,
    apply_mapping,
    build_quote_sheet,
    read_variant_table,
    split_constant_columns,
)

ok = True


def check(label, cond, detail=""):
    global ok
    ok &= bool(cond)
    print(("PASS  " if cond else "FAIL  ") + label + (f"  — {detail}" if detail and not cond else ""))


def family(n=6):
    """Standard is constant; size and qty vary; drawing is always blank."""
    return [
        {"Item code": f"STD-M{12 + i * 4}", "Description": "Threaded stud",
         "Standard": "DIN 976", "Size": f"M{12 + i * 4}", "Finish": "HDG 50 µm",
         "Qty — annual": str(100 + i * 20), "Drawing": ""}
        for i in range(n)
    ]


SPEC = ["Item code", "Description", "Standard", "Size", "Finish", "Qty — annual", "Drawing"]


def opened(data):
    return load_workbook(io.BytesIO(data))["Quote Sheet"]


def headers_of(ws):
    return [ws.cell(row=6, column=i).value for i in range(1, ws.max_column + 1)]


# ---- the variance rule -----------------------------------------------------
varying, constants = split_constant_columns(family(), SPEC)
check("a column that varies stays a column", "Size" in varying and "Qty — annual" in varying, str(varying))
check("a column identical on every row becomes a condition",
      "Standard" not in varying and any("DIN 976" in c for c in constants), str(constants))
check("the condition names the attribute, not just the value",
      any(c.startswith("Standard:") for c in constants), str(constants))
check("a column empty on every row is dropped entirely",
      "Drawing" not in varying and not any("Drawing" in c for c in constants), str(constants))
check("identity columns never collapse, even when constant",
      "Description" in varying, str(varying))

# A one-row family must not collapse into nothing but conditions.
v1, c1 = split_constant_columns(family(1), SPEC)
check("a single-SKU family keeps its identity columns",
      "Item code" in v1 and "Description" in v1, str(v1))


# ---- reading the customer's own workbook ----------------------------------
def customer_book(n_rows, header=("Part No", "Desc", "Matl", "Annual Qty")):
    wb = Workbook()
    ws = wb.active
    ws.append(["Acme Fasteners — enquiry sheet"])   # junk above the table
    ws.append([])
    ws.append(list(header))
    for i in range(n_rows):
        ws.append([f"P-{i:04d}", "Hex bolt", "8.8", 100 + i])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


t = read_variant_table(customer_book(12))
check("the header row is found below junk rows", t and t.header[:2] == ["Part No", "Desc"], str(t and t.header))
check("every body row is read", t and len(t.rows) == 12, str(t and len(t.rows)))

# The whole point: no MAX_EXCEL_ROWS truncation on this path.
t300 = read_variant_table(customer_book(300))
check("a 300-row BOM is read in full, not capped at 250",
      t300 and len(t300.rows) == 300, str(t300 and len(t300.rows)))

mapped = apply_mapping(t300, {"Item code": "Part No", "Description": "Desc",
                              "Material": "Matl", "Qty — annual": "Annual Qty"})
check("mapping renames the customer's columns to ours",
      mapped[0]["Item code"] == "P-0000" and mapped[0]["Qty — annual"] == "100", str(mapped[0]))
check("mapping preserves row order and count", len(mapped) == 300 and mapped[-1]["Item code"] == "P-0299")

missing = apply_mapping(t300, {"Item code": "Part No", "Finish": "Coating (not in their sheet)"})
check("a column they never supplied comes through empty, not as a crash",
      missing[0]["Finish"] == "" and missing[0]["Item code"] == "P-0000")

check("a workbook with nothing table-shaped returns None",
      read_variant_table(customer_book(0)) is None)


# ---- the workbook we produce ----------------------------------------------
data = build_quote_sheet(
    rfq_ref="WZ-TEST-1", title="Test Family", rows=family(6), spec_columns=SPEC,
    conditions=["50 µm hot dip galvanisation is acceptable."],
)
ws = opened(data)
hdr = headers_of(ws)

check("Sr leads the sheet", hdr[0] == "Sr", str(hdr))
check("the constant column is absent from the header", "Standard" not in hdr, str(hdr))
check("the varying columns are present", "Size" in hdr and "Qty — annual" in hdr, str(hdr))
# Finish is identical on every row in this fixture, so it belongs in the
# conditions block — not the header. Same rule as Standard.
check("a second constant column also collapses", "Finish" not in hdr, str(hdr))
check("every supplier column is present",
      all(q.name in hdr for q in DEFAULT_QUOTE_COLUMNS), str(hdr))
check("one grid row per SKU", ws.cell(row=12, column=1).value == 6 and ws.cell(row=13, column=1).value is None)

body = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
check("the passed condition appears once", body.count("50 µm hot dip galvanisation") == 1)
check("the derived condition appears too", "Standard: DIN 976" in body)
check("the constant value is not repeated down the grid", body.count("DIN 976") == 1, str(body.count("DIN 976")))

check("supplier cells are left empty for them to fill",
      ws.cell(row=7, column=len(hdr)).value is None)
check("the grid is frozen so identity stays visible", ws.freeze_panes == "D7", str(ws.freeze_panes))
check("the header row repeats when printed",
      str(ws.print_title_rows).replace("$", "") == "6:6", str(ws.print_title_rows))
check("an autofilter covers the grid", ws.auto_filter.ref.startswith("A6:"), str(ws.auto_filter.ref))

# Numbers must arrive as numbers, or the supplier cannot sort or total them.
qty_col = hdr.index("Qty — annual") + 1
check("a numeric column is written as a number, not text",
      isinstance(ws.cell(row=7, column=qty_col).value, (int, float)),
      repr(ws.cell(row=7, column=qty_col).value))
code_col = hdr.index("Item code") + 1
check("a part number is left as text", isinstance(ws.cell(row=7, column=code_col).value, str))


# ---- nothing the supplier must not see ------------------------------------
leaky = [dict(r, **{"Target price": "$1.20", "Customer": "Mumtaz Foods"}) for r in family(3)]
data = build_quote_sheet(rfq_ref="R", title="T", rows=leaky, spec_columns=SPEC)
body = "\n".join(str(c.value) for row in opened(data).iter_rows() for c in row if c.value)
check("a column not asked for never reaches the sheet",
      "Mumtaz" not in body and "1.20" not in body, body[:200])


# ---- custom supplier columns ----------------------------------------------
data = build_quote_sheet(
    rfq_ref="R", title="T", rows=family(3), spec_columns=SPEC,
    quote_columns=(QuoteColumn("Unit price\n(USD)", 11, "number", "$#,##0.000"),
                   QuoteColumn("Tooling cost\n(USD)", 12, "number", "$#,##0")),
)
hdr = headers_of(opened(data))
check("the supplier block is caller-defined, not fixed",
      "Tooling cost\n(USD)" in hdr and "MOQ\n(pcs)" not in hdr, str(hdr))


# ---- scale -----------------------------------------------------------------
big = [{"Item code": f"P-{i:04d}", "Description": "Hex bolt", "Size": f"M{8 + i % 40}",
        "Finish": "HDG", "Qty — annual": str(100 + i)} for i in range(300)]
t0 = time.perf_counter()
data = build_quote_sheet(rfq_ref="R", title="Big Family", rows=big,
                         spec_columns=["Item code", "Description", "Size", "Finish", "Qty — annual"])
elapsed = time.perf_counter() - t0
ws = opened(data)
check("300 SKUs all land in the grid", ws.cell(row=306, column=1).value == 300, str(ws.cell(row=306, column=1).value))
check(f"300 SKUs build quickly ({elapsed:.2f}s, {len(data)/1024:.0f} KB)", elapsed < 5.0, f"{elapsed:.2f}s")

# ---- filename --------------------------------------------------------------
from rfq_summary.quote_sheet import annexure_filename, build_family_annexures

check("named for the family, numbered for citation",
      annexure_filename(1, "Hex Bolts") == "Annexure 1 - Hex Bolts.xlsx",
      annexure_filename(1, "Hex Bolts"))
check("a second annexure counts on",
      annexure_filename(2, "Flat Washers").startswith("Annexure 2 - "))
check("characters OneDrive rejects are stripped",
      not set('":<>?/\\|*') & set(annexure_filename(1, 'Studs — M56 x 2000 / DIN 976 <rev B>')),
      annexure_filename(1, 'Studs — M56 x 2000 / DIN 976 <rev B>'))
check("the family name still survives that",
      "Studs" in annexure_filename(1, 'Studs / DIN 976') and
      "DIN 976" in annexure_filename(1, 'Studs / DIN 976'),
      annexure_filename(1, 'Studs / DIN 976'))
check("a trailing period is removed (OneDrive rejects it)",
      not annexure_filename(1, "Bolts.").replace(".xlsx", "").endswith("."),
      annexure_filename(1, "Bolts."))
check("an over-long name is trimmed, keeping the number",
      len(annexure_filename(3, "X" * 300)) < 110 and
      annexure_filename(3, "X" * 300).startswith("Annexure 3 - "))
check("a nameless family still yields a usable filename",
      annexure_filename(1, "") == "Annexure 1.xlsx", annexure_filename(1, ""))


# ---- one workbook per family, never one per RFQ ----------------------------
class _Annexure:
    def __init__(self, columns, rows, by_reference=False):
        self.columns, self.rows, self.by_reference = columns, rows, by_reference
        self.required = True


class _Product:
    def __init__(self, name, structure, annexure=None):
        self.name, self.structure, self.annexure = name, structure, annexure


COLS = ["variant_ref", "size", "finish", "quantity"]
VROWS = [[f"V{i}", f"M{8 + i * 2}", "HDG 50 µm", str(100 + i)] for i in range(8)]

built = build_family_annexures(rfq_ref="WZ-1", products=[
    _Product("Hex Bolts (family)", "family", _Annexure(COLS, VROWS)),
    _Product("Threaded Stud M56", "single"),                       # not a family
    _Product("Flat Washers (family)", "family", _Annexure(COLS, VROWS)),
    _Product("Empty Family", "family", _Annexure(COLS, [])),       # no rows
    _Product("Their Sheet", "family", _Annexure(COLS, VROWS, by_reference=True)),
])
names = [n for n, _ in built]
check("one workbook per family line", len(built) == 2, str(names))
check("numbered in order of the families found",
      names[0].startswith("Annexure 1 - Hex Bolts") and names[1].startswith("Annexure 2 - Flat Washers"),
      str(names))
check("a single line produces nothing", not any("Stud" in n for n in names), str(names))
check("a family with no rows produces nothing", not any("Empty" in n for n in names), str(names))
check("the customer's own workbook is not replaced by ours",
      not any("Their Sheet" in n for n in names), str(names))

ws = load_workbook(io.BytesIO(built[0][1]))["Quote Sheet"]
check("the sheet is titled for its family",
      "Hex Bolts" in str(ws.cell(row=1, column=1).value), str(ws.cell(row=1, column=1).value))
check("all eight variants are rows", ws.cell(row=14, column=1).value == 8)

# An RFQ of four single lines — the real Clint 7 shape — makes no file at all.
check("an RFQ with no families generates nothing",
      build_family_annexures(rfq_ref="WZ-2", products=[
          _Product("M56 Stud HDG", "single"), _Product("M56 Nut HDG", "single"),
          _Product("M56 Stud A4", "single"), _Product("M56 Nut A4", "single")]) == [])

# Dict-shaped variant rows are as valid as positional ones.
d = build_family_annexures(rfq_ref="WZ-3", products=[
    _Product("Dict Family", "family",
             _Annexure(COLS, [dict(zip(COLS, r)) for r in VROWS]))])
check("dict-shaped variant rows work too", len(d) == 1 and d[0][0].startswith("Annexure 1 - Dict"))

print("\nALL PASSED" if ok else "\nFAILURES ABOVE")
raise SystemExit(0 if ok else 1)
