"""
test_costing_workbook.py — the generated internal costing workbook.

The rules that matter: every generated tab says it is generated, nothing
the team owns (weight, rates) is filled in by us, a missing process rate never
blanks a total, the summary adds up against the source, and a template's
Quotation keeps working after its item table grows — lookups into Back-end
and the reference to A4 must not move, while everything below the table must.

Run:
    python test_costing_workbook.py
"""
import io
import sys
import zipfile

sys.path.insert(0, "src")

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.datavalidation import DataValidation

from rfq_summary.costing_workbook import (
    SUMMARY_TITLE, Commons, Line, Tab, assumed, build_costing_workbook, data, doubtful, needed, shift_formula,
)

ok = True


def check(label, cond, detail=""):
    global ok
    ok &= bool(cond)
    print(("PASS  " if cond else "FAIL  ") + label + (f"  — {detail}" if detail and not cond else ""))


def header_map(ws):
    return {ws.cell(2, c).value: openpyxl.utils.get_column_letter(c) for c in range(1, ws.max_column + 1)
            if ws.cell(2, c).value}


# ---- shift_formula ---------------------------------------------------------
check("a reference below the insert point moves", shift_formula("=VLOOKUP(A58,X1:Y2,2)", 25, 3).startswith("=VLOOKUP(A61"))
check("a reference above it stays", shift_formula("=IFERROR(VLOOKUP(A4, 'Back-end'!F2:K6, 6, FALSE), \"\")", 25, 3)
      == "=IFERROR(VLOOKUP(A4, 'Back-end'!F2:K6, 6, FALSE), \"\")")
check("another sheet's rows never move", "'Back-end'!B16:E29" in shift_formula("=VLOOKUP(A58,'Back-end'!B16:E29,2)", 10, 5))
check("absolute rows below still move", shift_formula("=$A$30+B30", 25, 2) == "=$A$32+B32")
check("whole-column references are untouched",
      shift_formula('=INDEX(E:E, MATCH("Freight", A:A, 0))', 1, 9) == '=INDEX(E:E, MATCH("Freight", A:A, 0))')
check("text that looks like a cell is untouched", shift_formula('="A30"&B30', 25, 1) == '="A30"&B31')
check("a plain value passes through", shift_formula("hello", 1, 5) == "hello")


# ---- a small RFQ -----------------------------------------------------------
def fastener_tab(n=3):
    lines = [Line(label=f"M{8 + i} Flat Washer", fields={"Stock code": data(f"0{i}"), "Description": data(f"M{8 + i} Flat Washer"),
                                                          "Standard": doubtful("DIN 9021") if i == 1 else data("DIN 125A")},
                  qty=data(1000 * (i + 1)), pack=data(100), pack_unit="pcs / box", rate_group="Flat washer",
                  remarks="Very low qty" if i == 2 else "")
             for i in range(n)]
    return Tab("Washers", "family", "weight_rate", [("Stock code", 9), ("Description", 30), ("Standard", 12)], lines,
               source="Washers.xlsx › Washers", source_lines=n + 1, skipped=[("Row 9", "Only a standard, no item")],
               rate_groups=["Flat washer"])


def fab_tab():
    cover = Line(label="Cover", fields={"Drawing no.": data("MT-1"), "Material": assumed("Stainless 1.4404 (assumed)")},
                 qty=needed(), qty_unit="nos", weight=data(4.62),
                 processes=[assumed("Cutting"), data("Welding"), assumed("Laser marking")])
    return Tab("Individual items", "individual", "process", [("Drawing no.", 14), ("Material", 18)], [cover],
               source="MT-1.pdf", source_lines=1)


book = openpyxl.load_workbook(io.BytesIO(build_costing_workbook([fastener_tab(), fab_tab()])))
check("every generated tab is marked (Zai)", all(n.startswith("(Zai) ") for n in book.sheetnames), str(book.sheetnames))
check("the summary comes first", book.sheetnames[0] == SUMMARY_TITLE, str(book.sheetnames))

ws = book["(Zai) Washers"]
h = header_map(ws)
check("the table starts at the top — no title rows", ws["A1"].value == "CUSTOMER DATA" and ws["A2"].value == "Stock code",
      f"{ws['A1'].value!r} / {ws['A2'].value!r}")
check("first product sits on row 3", ws["B3"].value == "M8 Flat Washer", str(ws["B3"].value))
check("qty and its unit are separate columns",
      ws[f"{h['Annual qty']}3"].value == 1000 and ws[f"{h['Qty unit']}3"].value == "pcs")
check("pack size and its unit too", ws[f"{h['Pack size']}3"].value == 100 and ws[f"{h['Pack unit']}3"].value == "pcs / box")
check("no qty-per-box / box-count duplicates", not any(k in h for k in ("Box Quantity", "Pcs in a Box", "Pcs")), str(list(h)))
check("weight is never filled in by us", ws[f"{h['Weight']}3"].value is None)
check("…and is orange so it is seen", ws[f"{h['Weight']}3"].fill.fgColor.rgb.endswith("FCD5B4"))
check("the doubtful standard is pink", ws[f"{h['Standard']}4"].font.color.rgb.endswith("E0218A"))
check("the rate column only points at the legend",
      str(ws[f"{h['Rate (INR / kg)']}3"].value).startswith("=IF($C$"), str(ws[f"{h['Rate (INR / kg)']}3"].value))
legend_rate = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Flat washer"]
check("the rate itself is an empty legend cell", legend_rate and ws.cell(legend_rate[0], 3).value is None)
check("missing inputs are spelled out", "Add weight" in str(ws[f"{h['Missing input']}3"].value))
fx = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "INR per GBP"]
check("shared assumptions live once, on the first tab", fx and ws.cell(fx[0], 3).value == 120)
check("remarks carried through", ws[f"{h['Remarks']}5"].value == "Very low qty")

fab = book["(Zai) Individual items"]
fh = header_map(fab)
check("eight process columns", sum(1 for k in fh if k.startswith("Process ")) == 8, str([k for k in fh if k.startswith("Process")]))
check("processes listed, black and red kept", fab[f"{fh['Process 2']}3"].value == "Welding"
      and fab[f"{fh['Process 1']}3"].font.color.rgb.endswith("C00000"))
names = [fab.cell(r, 2).value for r in range(1, fab.max_row + 1)]
check("a process the defaults lack is added to the legend", "Laser marking" in names)
mfg = str(fab[f"{fh['Mfg cost / pc (INR)']}3"].value)
check("a missing process rate does not blank the total",
      "COUNTIFS" not in mfg.split(',"",')[0], mfg)
check("…but is counted in Missing input", "process rate(s) missing" in str(fab[f"{fh['Missing input']}3"].value))
check("the stated weight is kept, black", fab[f"{fh['Weight']}3"].value == 4.62)
link = [r for r in range(1, fab.max_row + 1) if fab.cell(r, 2).value == "INR per GBP"]
check("other tabs link the shared assumptions", link and str(fab.cell(link[0], 3).value).startswith("='(Zai) Washers'!$C$"))
check("the process list can grow past the visible rows",
      "200" not in "" and any("ZaiProcesses" in n for n in book.defined_names))

s = book[SUMMARY_TITLE]
rows = {s.cell(r, 1).value: [s.cell(r, c).value for c in range(1, 11)] for r in range(2, 5)}
w = rows.get("(Zai) Washers")
check("summary: source vs extracted per tab", w and w[3] == 4 and w[4] == 3 and w[5] == 1, str(w))
check("summary: a gap is called out", w and "not extracted" in str(w[9]), str(w))
check("summary: doubtful and to-fill counts", w and w[6] == 1 and w[8] >= 3, str(w))
body = [s.cell(r, c).value for r in range(1, s.max_row + 1) for c in range(1, 4)]
check("summary lists what was skipped and why", "Only a standard, no item" in body)
check("summary indexes every line", "Cover" in body and "M10 Flat Washer" in body)


# ---- a template ------------------------------------------------------------
def template_bytes():
    from PIL import Image
    png = io.BytesIO()
    Image.new("RGB", (40, 20), "navy").save(png, "PNG")
    wb = openpyxl.Workbook()
    q = wb.active
    q.title = "Quotation"
    q["A4"] = "Our address"
    q["A12"], q["B12"] = "Item Name", "Price per 100 Pc"
    for r in range(13, 17):
        q[f"C{r}"] = f"=IF(LEN($A{r})>0,VLOOKUP($A{r},Costing!$B$1:$X$500,7,FALSE),\"\")"
    q["A17"], q["A18"] = "Freight", "Total DAP Price"
    q["D17"] = 25
    q["A20"] = "TERMS OF DELIVERY"
    q.merge_cells("A20:E20")
    q["A21"], q["B21"] = "Company Name", "=IFERROR(VLOOKUP(A4, 'Back-end'!F2:K6, 6, FALSE), \"\")"
    q.merge_cells("B21:E21")
    q["A24"], q["A25"] = "Paul", "=IFERROR(VLOOKUP(A24, 'Back-end'!B16:E29, 2, FALSE), \"\")"
    q.row_dimensions[21].height = 33
    dv = DataValidation(type="list", formula1='"DAP,EXW"')
    q.add_data_validation(dv)
    dv.add("B22:E22")
    for anchor in ("A1", "A23"):
        img = XLImage(io.BytesIO(png.getvalue()))
        q.add_image(img, anchor)
    wb.create_sheet("Back-end")["F2"] = "Our address"
    wb.create_sheet("Costing")["B7"] = "Sample"
    ex = wb.create_sheet("ExIm Insights")
    ex.add_image(XLImage(io.BytesIO(png.getvalue())), "A1")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


tab = fastener_tab(6)                      # 6 + 1 lines into a 4-row table: 3 rows inserted
raw = build_costing_workbook([tab, fab_tab()], template=template_bytes(), commons=Commons())
book = openpyxl.load_workbook(io.BytesIO(raw))
check("the template's own sample Costing tab is gone", "Costing" not in book.sheetnames, str(book.sheetnames))
check("generated tabs sit after Back-end, before the insights",
      book.sheetnames == ["Quotation", "Back-end", SUMMARY_TITLE, "(Zai) Washers", "(Zai) Individual items", "ExIm Insights"],
      str(book.sheetnames))
q = book["Quotation"]
check("every line is on the Quotation", [q[f"A{r}"].value for r in range(13, 20)] ==
      [f"M{8 + i} Flat Washer" for i in range(6)] + ["Cover"], str([q[f"A{r}"].value for r in range(13, 20)]))
check("price comes from the line's own tab", "'(Zai) Washers'!" in str(q["B13"].value), str(q["B13"].value))
check("an unknown qty shows blank, not 0", str(q["C19"].value).startswith('=IF('), str(q["C19"].value))
check("the old name lookup into Costing is gone", not any("Costing!" in str(c.value) for row in q.iter_rows() for c in row))
check("freight follows the table", q["A20"].value == "Freight" and q["D20"].value == 25, f"{q['A20'].value} {q['D20'].value}")
check("the total sums every line plus freight", q["D21"].value == "=SUM(D13:D19)+N(D20)", str(q["D21"].value))
check("terms moved down with their merge", q["A23"].value == "TERMS OF DELIVERY" and "A23:E23" in {str(m) for m in q.merged_cells.ranges})
check("the Back-end lookup and its A4 reference did not move",
      q["B24"].value == "=IFERROR(VLOOKUP(A4, 'Back-end'!F2:K6, 6, FALSE), \"\")", str(q["B24"].value))
check("a reference inside the moved block moved with it", "VLOOKUP(A27," in str(q["A28"].value), str(q["A28"].value))
check("row heights moved too", q.row_dimensions[24].height == 33)
dv_ranges = [str(d.sqref) for d in q.data_validations.dataValidation]
check("dropdowns below the table moved", "B25:E25" in dv_ranges, str(dv_ranges))
z = zipfile.ZipFile(io.BytesIO(raw))
drawings = [n for n in z.namelist() if n.startswith("xl/drawings/drawing") and n.endswith(".xml")]
check("pictures survive the round trip", len(drawings) == 2 and len([n for n in z.namelist() if "media" in n]) == 3,
      str(drawings))
import re
qd = re.sub(r"<(/?)[A-Za-z]\w*:", r"<\1", z.read("xl/drawings/drawing1.xml").decode())
check("a picture below the table moved with it", "<row>25</row>" in qd and "<row>0</row>" in qd
      and "<row>22</row>" not in qd,
      str(re.findall(r"<row>\d+</row>", qd)))

few = build_costing_workbook([fastener_tab(2)], template=template_bytes())
q = openpyxl.load_workbook(io.BytesIO(few))["Quotation"]
check("a short RFQ leaves the table in place", q["A17"].value == "Freight" and q["A14"].value == "M9 Flat Washer")
check("unused template rows are cleared", q["C15"].value is None and q["C16"].value is None)

print("\nALL PASSED" if ok else "\nFAILURES ABOVE")
raise SystemExit(0 if ok else 1)
